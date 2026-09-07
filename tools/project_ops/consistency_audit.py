"""
冻结数字一致性审计（P0-2）

来源：MathModeling-skills 仓库 frozen_numbers 机制，经 Modex 适配。
职责：论文中每个数字必须存在于 run_manifest.json，冲突即阻塞。

审计流程：
1. 从 run_manifest.json 提取所有冻结数字
2. 从论文 DOCX 提取所有数字
3. 比对：论文数字 ⊆ 冻结数字
4. 冲突项输出为阻塞性错误
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..docx.core.result_contract import load_manifest


@dataclass
class NumberConflict:
    """单个数字冲突记录"""
    number: str
    context: str
    source: str
    frozen_value: str | None = None
    severity: str = "error"  # error=阻塞, warning=提示


@dataclass
class ConsistencyAuditResult:
    """一致性审计结果"""
    total_numbers_in_paper: int = 0
    total_frozen_numbers: int = 0
    matched: int = 0
    conflicts: list[NumberConflict] = field(default_factory=list)
    warnings: list[NumberConflict] = field(default_factory=list)
    passed: bool = True

    @property
    def blocking_errors(self) -> list[NumberConflict]:
        return [c for c in self.conflicts if c.severity == "error"]

    def to_dict(self) -> dict[str, Any]:
        return {
            "total_numbers_in_paper": self.total_numbers_in_paper,
            "total_frozen_numbers": self.total_frozen_numbers,
            "matched": self.matched,
            "conflicts": [c.__dict__ for c in self.conflicts],
            "warnings": [c.__dict__ for c in self.warnings],
            "passed": self.passed,
            "blocking_errors": len(self.blocking_errors),
        }


# 数字提取正则：匹配整数、小数、科学计数法、百分比
# 用 ASCII lookbehind/lookahead 替代 \b，避免 % 与中文字符的边界冲突
_NUMBER_RE = re.compile(
    r"(?<![A-Za-z0-9])(\d+(?:\.\d+)?(?:[eE][+-]?\d+)?%?)(?![A-Za-z])"
)

# 排除模式：年份、页码、章节号等
_EXCLUDE_PATTERNS = [
    re.compile(r"^(19|20)\d{2}$"),  # 年份
    re.compile(r"^\d{1,3}$"),  # 纯数字页码/章节号
]


def _is_excluded_number(num_str: str) -> bool:
    """判断数字是否应排除（年份、页码等）"""
    for pattern in _EXCLUDE_PATTERNS:
        if pattern.match(num_str):
            return True
    return False


def extract_numbers_from_text(text: str) -> list[tuple[str, str]]:
    """
    从文本中提取数字及其上下文。

    Returns:
        [(number_str, context_str), ...]
    """
    results = []
    for match in _NUMBER_RE.finditer(text):
        num = match.group(1)
        if _is_excluded_number(num):
            continue
        start = max(0, match.start() - 30)
        end = min(len(text), match.end() + 30)
        context = text[start:end].replace("\n", " ")
        results.append((num, context))
    return results


def extract_numbers_from_docx(docx_path: Path) -> list[tuple[str, str]]:
    """从 DOCX 文件中提取所有数字"""
    try:
        from docx import Document
        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        # 也提取表格中的文本
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text
        return extract_numbers_from_text(full_text)
    except ImportError:
        raise RuntimeError("python-docx 未安装，无法提取 DOCX 数字")


def extract_frozen_numbers(manifest_path: Path) -> dict[str, str]:
    """
    从 run_manifest.json 提取所有冻结数字。

    Returns:
        {number_str: json_path, ...}
    """
    manifest = load_manifest(manifest_path)
    frozen: dict[str, str] = {}

    def _walk(obj: Any, path: str = "$") -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                _walk(v, f"{path}.{k}")
        elif isinstance(obj, list):
            for i, v in enumerate(obj):
                _walk(v, f"{path}[{i}]")
        elif isinstance(obj, (int, float)) and not isinstance(obj, bool):
            num_str = str(obj)
            frozen[num_str] = path

    _walk(manifest)
    return frozen


def audit_consistency(
    docx_path: Path,
    manifest_path: Path,
    *,
    tolerance: float = 1e-6,
) -> ConsistencyAuditResult:
    """
    执行冻结数字一致性审计。

    Args:
        docx_path: 论文 DOCX 路径
        manifest_path: run_manifest.json 路径
        tolerance: 浮点数比较容差

    Returns:
        ConsistencyAuditResult
    """
    result = ConsistencyAuditResult()

    # 提取论文数字
    paper_numbers = extract_numbers_from_docx(docx_path)
    result.total_numbers_in_paper = len(paper_numbers)

    # 提取冻结数字
    frozen = extract_frozen_numbers(manifest_path)
    result.total_frozen_numbers = len(frozen)

    # 比对
    for num_str, context in paper_numbers:
        # 尝试精确匹配
        if num_str in frozen:
            result.matched += 1
            continue

        # 尝试浮点数容差匹配
        try:
            paper_val = float(num_str.replace("%", ""))
            matched = False
            for frozen_str, path in frozen.items():
                try:
                    frozen_val = float(frozen_str)
                    if abs(paper_val - frozen_val) <= tolerance:
                        result.matched += 1
                        matched = True
                        break
                except ValueError:
                    continue
            if matched:
                continue
        except ValueError:
            pass

        # 未匹配
        conflict = NumberConflict(
            number=num_str,
            context=context,
            source="paper",
            frozen_value=None,
            severity="error",
        )
        result.conflicts.append(conflict)

    result.passed = len(result.blocking_errors) == 0
    return result


def audit_consistency_from_text(
    paper_text: str,
    manifest_path: Path,
    *,
    tolerance: float = 1e-6,
) -> ConsistencyAuditResult:
    """
    从纯文本执行一致性审计（用于非 DOCX 场景）。

    Args:
        paper_text: 论文全文文本
        manifest_path: run_manifest.json 路径
        tolerance: 浮点数比较容差

    Returns:
        ConsistencyAuditResult
    """
    result = ConsistencyAuditResult()

    paper_numbers = extract_numbers_from_text(paper_text)
    result.total_numbers_in_paper = len(paper_numbers)

    frozen = extract_frozen_numbers(manifest_path)
    result.total_frozen_numbers = len(frozen)

    for num_str, context in paper_numbers:
        if num_str in frozen:
            result.matched += 1
            continue

        try:
            paper_val = float(num_str.replace("%", ""))
            matched = False
            for frozen_str in frozen:
                try:
                    frozen_val = float(frozen_str)
                    if abs(paper_val - frozen_val) <= tolerance:
                        result.matched += 1
                        matched = True
                        break
                except ValueError:
                    continue
            if matched:
                continue
        except ValueError:
            pass

        result.conflicts.append(NumberConflict(
            number=num_str,
            context=context,
            source="paper",
            severity="error",
        ))

    result.passed = len(result.blocking_errors) == 0
    return result