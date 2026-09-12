import io
import json

from docx.oxml.ns import qn
import pathlib
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document

from tools.docx.core import paper_format
from tools.docx.core.structure_validation import _plot_pitfall_warnings
from tools.docx.core import paper_workflow
from tools.paper_search.scripts.hybrid_scholar import (
    CrossrefMetadata,
    HybridPaper,
    HybridScholar,
)
from tools.paper_search.scripts.openalex_scholar import Paper
from tools.pdf.scripts.convert_pdf_to_images import convert
from tools.xlsx.scripts.recalc import recalc


def test_hybrid_paper_is_json_serializable():
    paper = HybridPaper(
        title="A paper",
        authors=["A"],
        year=2025,
        citations=3,
        doi=None,
        abstract=None,
        sources=["openalex"],
    )

    payload = json.dumps(paper, default=HybridScholar._json_default)
    assert json.loads(payload)["citation_ready"] is False


def test_hybrid_search_rejects_nonpositive_limit():
    scholar = HybridScholar()

    with pytest.raises(ValueError, match="大于 0"):
        scholar.search_papers("test", limit=0)


def test_hybrid_search_crossref_verification_outputs_gbt7714(monkeypatch):
    openalex_paper = Paper(
        title="A robust optimization approach for the capacitated vehicle routing problem with demand uncertainty",
        authors=["Ilgaz Sungur", "Fernando Ordóñez", "Maged Dessouky"],
        publication_year=2008,
        cited_by_count=238,
        doi="10.1080/07408170701745378",
        abstract="abstract",
        venue="IIE Transactions",
        volume="40",
        issue="5",
        first_page="509",
        last_page="523",
        url="https://doi.org/10.1080/07408170701745378",
    )
    crossref_metadata = CrossrefMetadata(
        doi="10.1080/07408170701745378",
        title=openalex_paper.title,
        authors=["Ilgaz Sungur", "Fernando Ordóñez", "Maged Dessouky"],
        citation_authors=["Sungur I", "Ordóñez F", "Dessouky M"],
        author_families=["Sungur", "Ordóñez", "Dessouky"],
        year=2008,
        venue="IIE Transactions",
        volume="40",
        issue="5",
        pages="509-523",
        url="https://doi.org/10.1080/07408170701745378",
        work_type="journal-article",
    )
    scholar = HybridScholar()
    monkeypatch.setattr(
        scholar.openalex,
        "search_papers",
        lambda *args, **kwargs: [openalex_paper],
    )
    monkeypatch.setattr(scholar.crossref, "fetch", lambda doi: crossref_metadata)

    result = scholar.search_papers("robust optimization vehicle routing", limit=1)
    paper = result["verified"][0]

    assert result["stats"]["citation_ready"] == 1
    assert paper.citation_ready is True
    assert paper.sources == ["openalex", "crossref"]
    assert paper.verification_status == "crossref_verified"
    assert paper.citation_format == (
        "Sungur I, Ordóñez F, Dessouky M. "
        "A robust optimization approach for the capacitated vehicle routing "
        "problem with demand uncertainty[J]. IIE Transactions, 2008, 40(5): "
        "509-523. DOI: 10.1080/07408170701745378."
    )


def test_hybrid_search_keeps_mismatched_doi_candidate_unverified(monkeypatch):
    openalex_paper = Paper(
        title="A paper from OpenAlex",
        authors=["First Author"],
        publication_year=2024,
        cited_by_count=1,
        doi="10.1234/example",
        abstract=None,
    )
    crossref_metadata = CrossrefMetadata(
        doi="10.1234/example",
        title="A different paper",
        authors=["Other Author"],
        citation_authors=["Author O"],
        author_families=["Author"],
        year=2024,
        venue="Example Journal",
        volume="1",
        issue="1",
        pages="1-2",
        url="https://doi.org/10.1234/example",
        work_type="journal-article",
    )
    scholar = HybridScholar()
    monkeypatch.setattr(
        scholar.openalex,
        "search_papers",
        lambda *args, **kwargs: [openalex_paper],
    )
    monkeypatch.setattr(scholar.crossref, "fetch", lambda doi: crossref_metadata)

    result = scholar.search_papers("paper", limit=1)
    paper = result["papers"][0]

    assert result["verified"] == []
    assert paper.citation_ready is False
    assert paper.verification_status == "unverified"
    assert "题名不一致" in paper.verification_issues[0]


def test_hybrid_search_rejects_incomplete_citation_metadata(monkeypatch):
    openalex_paper = Paper(
        title="Incomplete metadata paper",
        authors=[],
        publication_year=None,
        cited_by_count=1,
        doi="10.1234/incomplete",
        abstract=None,
    )
    crossref_metadata = CrossrefMetadata(
        doi="10.1234/incomplete",
        title=openalex_paper.title,
        authors=[],
        citation_authors=[],
        author_families=[],
        year=None,
        venue=None,
        volume=None,
        issue=None,
        pages=None,
        url="https://doi.org/10.1234/incomplete",
        work_type="journal-article",
    )
    scholar = HybridScholar()
    monkeypatch.setattr(
        scholar.openalex,
        "search_papers",
        lambda *args, **kwargs: [openalex_paper],
    )
    monkeypatch.setattr(scholar.crossref, "fetch", lambda doi: crossref_metadata)

    result = scholar.search_papers("paper", limit=1)
    paper = result["papers"][0]
    payload = paper.to_dict()

    assert result["verified"] == []
    assert paper.citation_ready is False
    assert paper.citation_format is None
    assert payload["citation_format"] is None
    assert "书目信息不完整" in paper.verification_issues[0]


def test_hybrid_search_backfills_after_verification_failure(monkeypatch):
    bad_paper = Paper(
        title="Bad paper",
        authors=["First Author"],
        publication_year=2024,
        cited_by_count=100,
        doi="10.1234/bad",
        abstract=None,
    )
    good_paper = Paper(
        title="Good paper",
        authors=["Alice Smith"],
        publication_year=2024,
        cited_by_count=1,
        doi="10.1234/good",
        abstract=None,
    )
    good_metadata = CrossrefMetadata(
        doi="10.1234/good",
        title="Good paper",
        authors=["Alice Smith"],
        citation_authors=["Smith A"],
        author_families=["Smith"],
        year=2024,
        venue="Example Journal",
        volume="1",
        issue="1",
        pages="1-2",
        url="https://doi.org/10.1234/good",
        work_type="journal-article",
    )
    scholar = HybridScholar()
    monkeypatch.setattr(
        scholar.openalex,
        "search_papers",
        lambda *args, **kwargs: [bad_paper, good_paper],
    )
    monkeypatch.setattr(
        scholar.crossref,
        "fetch",
        lambda doi: (
            CrossrefMetadata(
                doi="10.1234/bad",
                title="Different paper",
                authors=["First Author"],
                citation_authors=["Author F"],
                author_families=["Author"],
                year=2024,
                venue="Example Journal",
                volume="1",
                issue="1",
                pages="1-2",
                url="https://doi.org/10.1234/bad",
                work_type="journal-article",
            )
            if doi.endswith("/bad")
            else good_metadata
        ),
    )

    result = scholar.search_papers("paper", limit=1)

    assert len(result["verified"]) == 1
    assert result["verified"][0].title == "Good paper"
    assert len(result["papers"]) == 2
    assert result["stats"]["verification_attempted"] == 2


def test_hybrid_search_formats_proceedings_as_conference(monkeypatch):
    openalex_paper = Paper(
        title="Conference paper",
        authors=["Alice Smith"],
        publication_year=2024,
        cited_by_count=1,
        doi="10.1234/conference",
        abstract=None,
    )
    crossref_metadata = CrossrefMetadata(
        doi="10.1234/conference",
        title="Conference paper",
        authors=["Alice Smith"],
        citation_authors=["Smith A"],
        author_families=["Smith"],
        year=2024,
        venue="Proceedings of Example",
        volume=None,
        issue=None,
        pages="1-2",
        url="https://doi.org/10.1234/conference",
        work_type="proceedings-article",
        event_location="Beijing",
        event_date="2024-07-01",
    )
    scholar = HybridScholar()
    monkeypatch.setattr(
        scholar.openalex,
        "search_papers",
        lambda *args, **kwargs: [openalex_paper],
    )
    monkeypatch.setattr(scholar.crossref, "fetch", lambda doi: crossref_metadata)

    result = scholar.search_papers("paper", limit=1)

    assert result["verified"][0].citation_ready is True
    assert result["verified"][0].citation_format == (
        "Smith A. Conference paper[C]//Proceedings of Example, "
        "Beijing, 2024-07-01: 1-2. DOI: 10.1234/conference."
    )


def test_hybrid_search_formats_book_as_monograph(monkeypatch):
    openalex_paper = Paper(
        title="A book on mathematical modeling",
        authors=["Alice Smith"],
        publication_year=2024,
        cited_by_count=1,
        doi="10.1234/book",
        abstract=None,
    )
    crossref_metadata = CrossrefMetadata(
        doi="10.1234/book",
        title="A book on mathematical modeling",
        authors=["Alice Smith"],
        citation_authors=["Smith A"],
        author_families=["Smith"],
        year=2024,
        venue=None,
        volume=None,
        issue=None,
        pages=None,
        url="https://doi.org/10.1234/book",
        work_type="book",
        publisher="Example Press",
        publisher_location="Beijing",
    )
    scholar = HybridScholar()
    monkeypatch.setattr(
        scholar.openalex,
        "search_papers",
        lambda *args, **kwargs: [openalex_paper],
    )
    monkeypatch.setattr(scholar.crossref, "fetch", lambda doi: crossref_metadata)

    result = scholar.search_papers("book mathematical modeling", limit=1)

    assert result["verified"][0].citation_format == (
        "Smith A. A book on mathematical modeling[M]. "
        "Beijing: Example Press, 2024. DOI: 10.1234/book."
    )


@pytest.mark.parametrize(
    ("work_type", "kwargs", "expected"),
    [
        (
            "dissertation",
            {
                "publisher": "Example University",
                "publisher_location": "Beijing",
            },
            "Smith A. A modeling dissertation[D]. Beijing: Example University, "
            "2024. DOI: 10.1234/reference.",
        ),
        (
            "report",
            {
                "publisher": "Example Research Institute",
                "publisher_location": "Shanghai",
            },
            "Smith A. A modeling report[R]. Shanghai: Example Research Institute, "
                "2024. DOI: 10.1234/reference.",
        ),
        (
            "book-chapter",
            {
                "venue": "Handbook of Mathematical Modeling",
                "publisher": "Example Press",
                "publisher_location": "Guangzhou",
            },
            "Smith A. A modeling chapter[M]//Handbook of Mathematical Modeling. "
            "Guangzhou: Example Press, 2024. DOI: 10.1234/reference.",
        ),
        (
            "standard",
            {
                "authors": [],
                "citation_authors": [],
                "standard_number": "GB/T 7714-2015",
                "publisher": "中国标准出版社",
                "publisher_location": "北京",
            },
            "GB/T 7714-2015. A modeling standard[S]. 北京: 中国标准出版社, "
            "2024. DOI: 10.1234/reference.",
        ),
        (
            "patent",
            {
                "patent_country": "中国",
                "patent_number": "CN123456789A",
                "published_date": "2024-01-02",
                "url": "https://example.test/patent",
            },
            "Smith A. A modeling patent: 中国, CN123456789A[P]. 2024-01-02. "
            "https://example.test/patent.",
        ),
        (
            "dataset",
            {
                "url": "https://example.test/dataset",
            },
            "Smith A. A modeling dataset[DS/OL]. 2024. "
            "https://example.test/dataset.",
        ),
        (
            "posted-content",
            {
                "url": "https://example.test/preprint",
            },
            "Smith A. A modeling preprint[EB/OL]. 2024. "
            "https://example.test/preprint.",
        ),
    ],
)
def test_hybrid_paper_formats_nonjournal_reference_types(
    work_type, kwargs, expected
):
    data = dict(kwargs)
    paper = HybridPaper(
        title={
            "dissertation": "A modeling dissertation",
            "report": "A modeling report",
            "standard": "A modeling standard",
            "patent": "A modeling patent",
            "dataset": "A modeling dataset",
            "posted-content": "A modeling preprint",
            "book-chapter": "A modeling chapter",
        }[work_type],
        authors=data.pop("authors", ["Alice Smith"]),
        year=2024,
        citations=0,
        doi="10.1234/reference",
        abstract=None,
        citation_authors=data.pop("citation_authors", ["Smith A"]),
        verification_status="crossref_verified",
        work_type=work_type,
        **data,
    )

    assert paper.citation_ready is True
    assert paper.citation_format == expected


def test_emit_progress_calculates_snapshot_once(monkeypatch):
    calls = []

    def fake_snapshot(doc, stage, rendered_pages):
        calls.append((doc, stage, rendered_pages))
        return {"stage": stage}

    monkeypatch.setattr(paper_workflow, "progress_snapshot", fake_snapshot)
    stream = io.StringIO()

    result = paper_workflow.emit_progress("doc", "validated", 2, stream)

    assert result == {"stage": "validated"}
    assert calls == [("doc", "validated", 2)]
    assert json.loads(stream.getvalue()) == {"stage": "validated"}


def test_pdf_conversion_creates_output_directory(tmp_path: Path):
    import fitz

    source = tmp_path / "source.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(source)

    output = tmp_path / "nested" / "images"
    convert(source, output, max_dim=300, dpi=72)

    assert (output / "page_1.png").is_file()


def test_recalc_rejects_non_positive_timeout(tmp_path: Path):
    workbook = tmp_path / "book.xlsx"
    workbook.write_bytes(b"not used")

    assert "超时秒数" in recalc(workbook, timeout=0)["error"]


def test_paper_quality_gate_rejects_short_rendered_paper():
    doc = paper_format.new_document()
    paper_format.title(doc, "论文题目")
    paper_format.abstract_title(doc)
    paper_format.body(doc,
        '本文针对某类带约束的组合优化问题建立了混合整数规划模型并设计分支定界求解算法，旨在为实际调度决策提供定量化工具与理论支撑。'
        '通过理论分析与数值实验，验证了模型的有效性和算法的有限步收敛性，为后续研究奠定基础。'
        '主要结果包括：目标函数值较经典贪婪基准降低15%以上，平均计算时间控制在秒级范围内，满足实时调度需求。'
        '灵敏度分析表明模型关键参数在正负20%扰动区间内解的质量稳健，结论对实际工程调度问题具有参考价值。'
        '模型考虑了决策变量的整数约束与线性不等式约束的平衡，求解过程采用深度优先搜索与最优界剪枝策略。'
        '实验数据来源于公开标准测试集，经过严格预处理后用于模型验证和参数标定，覆盖小规模到中规模共三组算例。'
        '对比实验还检验了算法在不同初始条件下的表现，结果表明所提方法在中等规模算例上表现优异，收敛率超过90%。'
        '未来工作可将框架扩展至多目标情形或引入启发式加速策略以处理更大规模实例，并探索与机器学习方法的深度结合。')
    paper_format.keywords(doc, "优化；预测；模型；分析")
    paper_format.heading1(doc, "一、问题重述")

    issues = paper_format.validate_paper_structure(
        doc,
        rendered_pages=paper_format.RenderedPageCount(8, 7),
        require_rendered_pages=True,
        project_root=None,
    )

    assert any("当前总页数 8 页" in issue for issue in issues)
    assert any("DOCX估算篇幅约" in issue for issue in issues)


def test_project_audit_hub_references_resolve():
    from tools.project_ops.project_audit import hub_reference_errors

    assert hub_reference_errors() == []


def test_self_check_docx_option_reports_template_tone(tmp_path: Path, capsys):
    doc = paper_format.new_document()
    paper_format.title(doc, "论文题目")
    paper_format.abstract_title(doc)
    paper_format.body(doc,
        '本文针对某优化问题建立了数学模型并设计求解算法。'
        '首先，此模型具有重要意义，并为后续求解奠定坚实基础。'
        '通过理论分析与数值实验验证了模型的有效性，主要结果包括目标函数值较基准降低15%以上。'
        '灵敏度分析表明模型参数在合理区间内稳健，结论对实际问题具有参考价值。'
        '模型考虑了约束条件的平衡，求解过程采用迭代优化策略。'
        '实验数据来源于公开数据集，经过预处理后用于模型验证和参数标定。'
        '对比实验检验了算法在不同条件下的表现，结果表明所提方法有效且稳健。'
        '未来工作可将框架扩展至多目标情形或引入启发式策略以处理更大规模实例。'
        '此外，本文还讨论了模型的适用范围与局限性，并提出了改进方向。'
        '研究过程中采用了文献调研、理论推导、数值仿真相结合的方法论。'
        '最终形成的求解框架具有较好的通用性，可迁移至类似优化问题。'
        '关键词选取覆盖了问题建模、求解方法与结果分析的核心概念。'
        '整体研究遵循从问题分析到模型建立再到求解验证的完整流程，具有系统性。')
    paper_format.keywords(doc, "优化；预测；模型；分析")
    docx_path = tmp_path / "paper.docx"
    doc.save(docx_path)

    from tools.docx.scripts import self_check

    assert self_check.main(["--docx", str(docx_path)]) == 1
    captured = capsys.readouterr()
    assert "疑似 AI 味" in captured.out
    assert "模板化表达扫描失败" in captured.err


def _ref_paragraph(text):
    class _P:
        def __init__(self, value):
            self.text = value
    return _P(text)


def test_reference_year_gate_rejects_pre_2016():
    from tools.docx.core.structure_validation import _reference_issues

    paragraphs = [
        _ref_paragraph("参考文献"),
        _ref_paragraph("[1] 张三. 旧方法[J]. 老学报, 2015, 3(2): 1-9."),
        _ref_paragraph("[2] 李四. 新方法[J]. 学报, 2018(4): 12-20."),
        _ref_paragraph("[3] 王五. 无年份著作[M]. 北京: 出版社."),
    ]
    issues = _reference_issues(paragraphs)
    assert any("年份 2015" in issue for issue in issues)
    assert not any("2018" in issue and "年份" in issue for issue in issues)
    assert not any("无年份著作" in issue for issue in issues)


def test_body_style_multiple_125_spacing_and_word_rerender_guards():
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn

    doc = paper_format.new_document()
    normal = doc.styles[paper_format.BODY_STYLE]
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert abs(normal.paragraph_format.line_spacing - 1.25) < 1e-6
    assert normal.paragraph_format.first_line_indent == paper_format.Pt(24)

    doc_defaults = doc.styles.element.find(qn("w:docDefaults"))
    assert doc_defaults is not None
    ppr_default = doc_defaults.find(qn("w:pPrDefault"))
    assert ppr_default is not None
    snap = ppr_default.find(qn("w:pPr")).find(qn("w:snapToGrid"))
    assert snap is not None and snap.get(qn("w:val")) == "0"
    for name in (paper_format.BODY_STYLE, paper_format.HEADING1_STYLE):
        style_snap = doc.styles[name]._element.get_or_add_pPr().find(qn("w:snapToGrid"))
        assert style_snap is not None and style_snap.get(qn("w:val")) == "0"
    assert all(
        len(section._sectPr.findall(qn("w:docGrid"))) == 0
        for section in doc.sections
    )


def test_export_latex_source_covers_core_elements(tmp_path):
    """LaTeX 源码版包含标题/章节/公式/关键词/表格/题注，公式体不转义。"""
    from tools.docx.core.latex_export import export_latex_source

    doc = paper_format.new_document()
    paper_format.title(doc, "基于测试模型的论文")
    paper_format.abstract_title(doc)
    paper_format.body(doc,
        '本文针对某优化问题建立了数学模型并设计求解算法。'
        '通过理论分析与数值实验验证了模型的有效性，主要结果包括目标函数值较基准降低15%以上。'
        '灵敏度分析表明模型参数在合理区间内稳健，结论对实际问题具有参考价值。'
        '模型考虑了约束条件的平衡，求解过程采用迭代优化策略。'
        '实验数据来源于公开数据集，经过预处理后用于模型验证和参数标定。'
        '对比实验检验了算法在不同条件下的表现，结果表明所提方法有效且稳健。'
        '未来工作可将框架扩展至多目标情形或引入启发式策略以处理更大规模实例。'
        '此外，本文还讨论了模型的适用范围与局限性，并提出了改进方向。'
        '研究过程中采用了文献调研、理论推导、数值仿真相结合的方法论。'
        '最终形成的求解框架具有较好的通用性，可迁移至类似优化问题。'
        '关键词选取覆盖了问题建模、求解方法与结果分析的核心概念。'
        '整体研究遵循从问题分析到模型建立再到求解验证的完整流程，具有系统性。'
        '模型建立过程中充分考虑了实际约束条件与目标函数之间的平衡关系。')
    paper_format.keywords(doc, "优化；预测；模型；分析")
    paper_format.heading1(doc, "一、问题重述")
    paper_format.body(doc, "对问题的重述文字。")
    paper_format.heading1(doc, "五、模型建立与求解")
    paper_format.equation(doc, "y = k x + b", number="(1)")
    paper_format.body(doc, "其中 k 为斜率。")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "方案"
    table.rows[0].cells[1].text = "误差"
    table.rows[1].cells[0].text = "A"
    table.rows[1].cells[1].text = "0.12"

    out = export_latex_source(doc, tmp_path / "完整论文.tex")
    text = out.read_text(encoding="utf-8")
    assert "fontset=fandol" in text  # 字体集确定性：Overleaf/本地 TeX Live 开箱可编译
    assert r"\section{一、问题重述}" in text
    assert r"\begin{equation}y = k x + b\tag{1}\end{equation}" in text
    assert r"\noindent\textbf{关键词：}" in text
    assert r"\begin{tabularx}{\textwidth}{cX}" in text and r"\toprule" in text
    assert r"\end{document}" in text


def test_save_document_writes_latex_source_first(tmp_path, monkeypatch, capsys):
    """save_document 在 DOCX 交付前先写出完整论文.tex（同一内容快照）。"""
    import tools.docx.core.structure_validation as sv

    monkeypatch.setattr(sv, "validate_paper_structure", lambda *a, **k: [])
    doc = paper_format.new_document()
    paper_format.title(doc, "论文题目")
    paper_format.body(doc, "正文内容。")
    out = paper_format.save_document(doc, tmp_path, overwrite=True)
    assert out.exists()
    tex = tmp_path / "完整论文.tex"
    assert tex.exists()
    assert "论文题目" in tex.read_text(encoding="utf-8")
    assert r"\end{document}" in tex.read_text(encoding="utf-8")
    err_lines = [l for l in capsys.readouterr().err.splitlines() if l.strip()]
    assert len(err_lines) == 1
    payload = json.loads(err_lines[0])
    assert payload["stage"] == "delivered" and payload["path"].endswith("完整论文.docx")


def test_save_document_tex_publish_is_atomic(tmp_path, monkeypatch):
    """tex 落位用 os.replace 原子替换：replace 失败时旧 tex 完整保留、无 .tmp 残留。"""
    import os
    import tools.docx.core.structure_validation as sv

    monkeypatch.setattr(sv, "validate_paper_structure", lambda *a, **k: [])
    doc = paper_format.new_document()
    paper_format.title(doc, "论文题目")
    paper_format.body(doc, "正文内容。")
    paper_format.save_document(doc, tmp_path, overwrite=True)
    tex = tmp_path / "完整论文.tex"
    old_content = tex.read_text(encoding="utf-8")

    def _blocked(src, dst):
        raise PermissionError("目标被占用")
    docx_before = (tmp_path / "完整论文.docx").read_bytes()
    monkeypatch.setattr(os, "replace", _blocked)
    with pytest.raises(PermissionError):
        paper_format.save_document(doc, tmp_path, overwrite=True)
    assert tex.read_text(encoding="utf-8") == old_content  # 旧 tex 完整
    assert (tmp_path / "完整论文.docx").read_bytes() == docx_before  # tex 先落位失败 → docx 不动
    assert not list(tmp_path.glob("*.tmp"))  # 无临时残留


def test_reference_year_gate_ignores_doi_and_page_digits():
    """DOI/页码/编号里的 4 位数字不得触发年份门禁（2016 后文献不误杀）。"""
    from tools.docx.core.structure_validation import _reference_issues

    paragraphs = [
        _ref_paragraph("参考文献"),
        _ref_paragraph("[4] Smith A. A method[J]. Journal, 2018, 12(3): 2015-2020. DOI: 10.1234/1999.abc."),
        _ref_paragraph("[5] Doe J. Another work[J]. 学报, 2021(2): 17-25."),
    ]
    issues = _reference_issues(paragraphs)
    assert not any("年份" in issue and ("[4]" in issue or "[5]" in issue) for issue in issues)


def test_clipped_object_gate_resolves_style_chain_spacing():
    """样式级为多倍 1.25 时，段级被手工改回固定 18 磅且含图片 → 直接格式即判出并拒存。"""
    import base64 as _b64

    from docx.shared import Pt

    doc = paper_format.new_document()
    p = doc.add_paragraph()  # Normal 样式：样式级多倍 1.25；随后段级改回固定值模拟手工重排
    png_1px = _b64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
        "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )
    p.add_run().add_picture(io.BytesIO(png_1px))
    p.paragraph_format.line_spacing = Pt(18)
    p.paragraph_format.line_spacing_rule = None if False else p.paragraph_format.line_spacing_rule
    issues = paper_format._clipped_object_issues(doc)
    assert len(issues) == 1 and "图片" in issues[0]


def test_clipped_object_gate_clear_body_text_ok():
    """正文纯文本段落（多倍 1.25 行距）不触发裁剪门禁。"""
    doc = paper_format.new_document()
    paper_format.body(doc, "普通正文段落，不含对象。")
    assert paper_format._clipped_object_issues(doc) == []


# ===== 瘦身白名单清理与 files/ 保护区回归 =====

def _mk_project(tmp_path):
    code = tmp_path / "code"
    data = tmp_path / "results" / "数据"
    files = tmp_path / "files"
    for d in (code, data, files):
        d.mkdir(parents=True, exist_ok=True)
    return code, data, files


def test_cleanup_never_touches_files_dir_or_project_root(tmp_path):
    """files/（赛题原件）与项目根层文件永不进入清理候选。"""
    from tools.project_ops.project_cleanup import plan_cleanup
    code, data, files = _mk_project(tmp_path)
    (files / "赛题.pdf").write_bytes(b"%PDF-1.4")
    (files / "原附录.xlsx").write_bytes(b"PK")
    (tmp_path / "我的笔记.txt").write_text("用户文件", encoding="utf-8")

    targets, _ = plan_cleanup(tmp_path)
    assert all("files" not in p.parts for p in targets)
    assert all(p.name != "我的笔记.txt" for p in targets)


def test_cleanup_preserves_paper_work_dir(tmp_path):
    """.paper_work 草稿目录不强制删除，保留供回溯。"""
    from tools.project_ops.project_cleanup import plan_cleanup
    _mk_project(tmp_path)
    pw = tmp_path / ".paper_work"
    pw.mkdir()
    (pw / "01_abstract.md").write_text("x", encoding="utf-8")

    targets, _ = plan_cleanup(tmp_path)
    assert pw not in targets


def test_plot_pitfall_warnings_flags_pie_twinx_jet(tmp_path):
    """W6 画图坑预警：饼图/双Y轴/jet 色图触发 W 级预警，干净脚本不触发。"""
    code = tmp_path / "code"
    code.mkdir()
    (code / "viz.py").write_text(
        "import matplotlib.pyplot as plt\n"
        "plt.pie([1, 2])\n"
        "ax = plt.gca().twinx()\n"
        "cmap = 'jet'\n",
        encoding="utf-8",
    )
    ws = _plot_pitfall_warnings(tmp_path)
    assert len(ws) == 3 and all("画图避坑清单" in w for w in ws)
    (code / "clean.py").write_text(
        "import matplotlib.pyplot as plt\nplt.plot([1, 2], cmap_v='viridis')\n", encoding="utf-8"
    )
    ws2 = _plot_pitfall_warnings(tmp_path)
    assert len(ws2) == 3 and all("viz.py" in w for w in ws2)


# ===== 新版契约补充覆盖 =====


def test_nine_step_verification_no_report_by_default(tmp_path):
    """9 步验收默认不落盘（瘦身）；显式 write_report=True 才写报告。"""
    from tools.project_ops.nine_step_verification import run_verification
    (tmp_path / "results").mkdir()
    run_verification(tmp_path, paper_text="一、问题重述\n测试正文。")
    assert not (tmp_path / "results" / "论文验收报告.md").exists()
    run_verification(tmp_path, paper_text="一、问题重述\n测试正文。", write_report=True)
    assert (tmp_path / "results" / "论文验收报告.md").exists()


def test_soft_doc_structure_markers():
    """软文档结构标记：去AI味含 Humanizer 3.0 Core + Math Extension；算法资料 7 卡；建模通用规范=防错速查；模板并入设计原则。"""
    root = pathlib.Path(__file__).resolve().parents[1]
    deai = (root / "文档/去AI味指南.md").read_text(encoding="utf-8")
    assert "## 九、Humanizer 3.0 Core" in deai and "## 十、Modex Math Extension" in deai
    algo_cards = list((root / "知识库/算法资料").glob("*.md"))
    assert len(algo_cards) == 7
    assert all("选型卡" in f.read_text(encoding="utf-8") for f in algo_cards)
    general = (root / "知识库/建模通用规范.md").read_text(encoding="utf-8")
    assert "题型防错速查" in general and "2analysis-modeling" not in general
    design = (root / "知识库/方法库/设计原则.md").read_text(encoding="utf-8")
    assert "分层方法卡模板" in design


def test_tournament_operational_section():
    """Tournament 操作细则：冠军挑战者文档含实验记录格式、评分公式与预算约束。"""
    root = pathlib.Path(__file__).resolve().parents[1]
    cc = (root / "知识库/建模增强/冠军挑战者建模流程.md").read_text(encoding="utf-8")
    assert "## 11. Tournament 操作细则" in cc
    assert "tournament_log.md" in cc
    assert "S(M) = w_1" in cc or "S(M)=w1" in cc
    assert "预算约束" in cc


def test_data_file_warnings_flag_bom_and_extra_json(tmp_path):
    """W8 数据文件格式：csv 缺 UTF-8-SIG BOM、数据目录多余 json → 预警；合规文件不触发。"""
    from tools.docx.core.structure_validation import _data_file_warnings
    data = tmp_path / "results" / "数据"
    data.mkdir(parents=True)
    (data / "no_bom.csv").write_text("a,b\n1,2\n", encoding="utf-8")  # 无 BOM
    (data / "good.csv").write_text("a,b\n1,2\n", encoding="utf-8-sig")  # 带 BOM
    (data / "spss_outputs.csv").write_text("name,unit\n", encoding="utf-8-sig")  # 工具链登记 csv
    (data / "result.json").write_text("{}", encoding="utf-8")  # 解题数据 json → 预警
    ws = _data_file_warnings(str(tmp_path))
    assert len(ws) == 2
    assert any(w.startswith("结果 CSV") and "no_bom.csv" in w for w in ws)
    assert any("不得使用 JSON" in w and "result.json" in w for w in ws)


def test_plot_pitfall_warnings_flag_bare_legend_and_best(tmp_path):
    """W6 扩展：裸 legend() 与 loc='best' 触发 P20 图例遮挡预警。"""
    from tools.docx.core.structure_validation import _plot_pitfall_warnings
    code = tmp_path / "code"
    code.mkdir()
    (code / "fig.py").write_text(
        "import matplotlib.pyplot as plt\n"
        "plt.legend()\n"
        "ax.legend(loc='best')\n",
        encoding="utf-8",
    )
    ws = _plot_pitfall_warnings(tmp_path)
    assert sum("P20" in w for w in ws) >= 2


def test_figure_table_context_warnings(tmp_path):
    """W9 图表上下文：前无引导/后无解释 → 预警；规范写法（前引出后解释）不触发。"""
    from docx.enum.style import WD_STYLE_TYPE
    from tools.docx.core.paper_format import CAPTION_STYLE
    from tools.docx.core.structure_validation import _figure_table_context_warnings

    def build(lead, explain):
        doc = Document()
        try:
            doc.styles.add_style(CAPTION_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        except (KeyError, ValueError):
            pass
        doc.add_paragraph("一、问题重述")
        if lead:
            doc.add_paragraph(lead)
        doc.add_paragraph("图1 测试对比", style=CAPTION_STYLE)
        if explain:
            doc.add_paragraph(explain)
        doc.add_paragraph("表1 符号说明", style=CAPTION_STYLE)
        tb = doc.add_table(rows=2, cols=2)
        tb.rows[0].cells[0].text = "符号"
        tb.rows[1].cells[0].text = "x"
        if explain:
            doc.add_paragraph("表1 给出全文符号体系，含义与单位逐列对应。")
        return doc

    bad = build(None, None)
    ws = _figure_table_context_warnings(bad)
    assert any("图1" in w and "后置解释" in w for w in ws)
    assert any("表1" in w and "前置引导" in w for w in ws)
    assert any("表1" in w and "后置解释" in w for w in ws)

    good = build("图1 展示两种方案的误差对比结果。",
                 "图1 中方案 A 在前 10 轮误差下降最快，原因是学习率设置更保守。")
    assert _figure_table_context_warnings(good) == []

def test_figure_table_lead_in_warnings(tmp_path):
    """W9 图表引出：紧跟标题/连续图表/紧跟标题后首图 → 预警；正常引出→解释不触发。"""
    from tools.docx.core.structure_validation import _figure_table_lead_in_warnings
    doc = Document()
    doc.add_paragraph("一、问题重述")
    doc.add_paragraph("图1 结果对比")            # 标题后直接图题 → 缺引出
    doc.add_paragraph("图2 灵敏度曲线")          # 连续图表 → 缺引出
    doc.add_paragraph("表1 主要符号说明")        # 仍连续 → 缺引出
    doc.add_paragraph("为评估模型稳健性，图3 给出扰动下的目标函数变化。")  # 引出
    doc.add_paragraph("图3 灵敏度分析")
    doc.add_paragraph("图3 显示扰动 ±20% 内目标函数波动小于 2%，模型稳健。")     # 解释
    doc.add_paragraph("为对比各方案的优劣，表2 汇总了关键指标。")                             # 引出
    doc.add_paragraph("表2 方案对比")
    issues = _figure_table_lead_in_warnings(doc)
    assert len(issues) == 3 and all("缺引出" in i for i in issues)


def test_rebuild_extracts_title_and_ai_declaration_heading():
    """rebuild 提炼器：首个文本段 → title；AI工具使用声明 → heading1（防标题变 body）。"""
    from docx import Document as D
    from tools.docx.core.paper_workflow import _extract_blocks
    doc = D()
    doc.add_paragraph("基于测试模型的论文")
    doc.add_paragraph("一、问题重述")
    doc.add_paragraph("正文。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("AI工具使用声明")
    blocks, _ = _extract_blocks(doc)
    kinds = [b["kind"] for b in blocks]
    assert kinds[0] == "title"
    assert blocks[-1]["kind"] == "heading1" and blocks[-1]["text"] == "AI工具使用声明"


def test_export_latex_uses_original_image_filename(tmp_path):
    """tex 的 includegraphics 用 results/图片/ 的原始生成文件名（哈希回查），非 DOCX 内部部件名。"""
    import base64 as b64

    from tools.docx.core.latex_export import export_latex_source

    png = b64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
        "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )
    img_dir = tmp_path / "results" / "图片"
    img_dir.mkdir(parents=True)
    (img_dir / "2_Q1_灵敏度分析.png").write_bytes(png)

    doc = paper_format.new_document()
    paper_format.heading1(doc, "六、模型检验与分析")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(png))
    paper_format.figure_caption(doc, "图1 灵敏度分析")

    out = export_latex_source(doc, tmp_path / "完整论文.tex", project_root=tmp_path)
    tex = out.read_text(encoding="utf-8")
    assert "2_Q1_灵敏度分析.png" in tex
    assert "image1.png" not in tex


def test_claim_strength_guard():
    """Claim Guard：检测结论强度升级（无证据强断言）。"""
    from tools.docx.core.structure_validation import _claim_strength_issues
    from docx import Document

    doc = Document()
    doc.add_paragraph('参数扰动±20%时目标函数波动<2.1%')
    assert _claim_strength_issues(doc) == []

    doc.add_paragraph('实验证明该模型具有普适性')
    issues = _claim_strength_issues(doc)
    assert len(issues) == 1
    assert '普适性' in issues[0]

    doc2 = Document()
    doc2.add_paragraph('模型鲁棒性强')
    issues2 = _claim_strength_issues(doc2)
    assert len(issues2) == 1
    assert '鲁棒性' in issues2[0]

    doc3 = Document()
    doc3.add_paragraph('该方法显著优于传统模型')
    issues3 = _claim_strength_issues(doc3)
    assert len(issues3) == 1
    assert '显著优于' in issues3[0]


def test_keyword_body_consistency():
    """关键词须在正文中出现。"""
    from tools.docx.core.structure_validation import _keyword_body_consistency_issues
    from docx import Document

    doc = Document()
    doc.add_paragraph('关键词：多目标优化；供应链；遗传算法')
    doc.add_paragraph('一、问题重述')
    doc.add_paragraph('本文研究供应链中的多目标优化问题，使用遗传算法求解。')
    doc.add_paragraph('参考文献')
    assert _keyword_body_consistency_issues(doc) == []

    doc2 = Document()
    doc2.add_paragraph('关键词：多目标优化；量子计算；遗传算法')
    doc2.add_paragraph('一、问题重述')
    doc2.add_paragraph('本文研究多目标优化问题，使用遗传算法求解。')
    doc2.add_paragraph('参考文献')
    issues = _keyword_body_consistency_issues(doc2)
    assert len(issues) == 1
    assert '量子计算' in issues[0]


def test_subquestion_completeness():
    """每问须具备变量设定、公式、求解/结果。"""
    from tools.docx.core.structure_validation import _subquestion_completeness_issues
    from docx import Document
    from lxml import etree

    OMML_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    OMML = f'<m:oMath xmlns:m="{OMML_NS}"><m:r><m:t>x</m:t></m:r></m:oMath>'

    def add_omath(para):
        para._p.append(etree.fromstring(OMML))

    doc = Document()
    doc.add_paragraph('五、模型建立与求解')
    doc.add_paragraph('5.1 问题一')
    doc.add_paragraph('设决策变量为x，目标函数为最小化成本。')
    add_omath(doc.add_paragraph(''))
    doc.add_paragraph('代入数据求解得到最优解x=3。')
    doc.add_paragraph('5.2 问题二')
    doc.add_paragraph('令y为时间变量，建立如下模型。')
    add_omath(doc.add_paragraph(''))
    doc.add_paragraph('计算结果如图1所示。')
    doc.add_paragraph('参考文献')
    assert _subquestion_completeness_issues(doc) == []

    doc2 = Document()
    doc2.add_paragraph('五、模型建立与求解')
    doc2.add_paragraph('5.1 问题一')
    doc2.add_paragraph('这个问题我们进行了分析。')
    doc2.add_paragraph('参考文献')
    issues = _subquestion_completeness_issues(doc2)
    assert len(issues) == 1
    assert '5.1' in issues[0]
    assert '变量设定' in issues[0]
    assert '数学公式' in issues[0]
    assert '求解' in issues[0]


def test_model_eval_gate_requires_pros5_cons4_and_no_prose():
    """模型评价闸门：优点 ≥5、局限与改进 ≥4 逐条编号；非编号大段拒存。"""
    from tools.docx.core.structure_validation import _model_eval_bullet_format_issues

    doc = Document()
    paras = [
        '七、模型评价与改进',
        '7.1 优点',
    ] + [f'（{i}）优点条目 {i}：约束完整且与基准对照误差 0.9%，属于本题可验证事实。'
         for i in range(1, 6)] + [
        '7.2 局限与改进方向',
    ] + [f'（{i}）改进条目 {i}：将参数改为随机变量并用鲁棒优化重解。'
         for i in range(1, 5)]
    for tx in paras:
        doc.add_paragraph(tx)
    assert _model_eval_bullet_format_issues(doc) == []


def test_model_eval_gate_rejects_few_items_and_stray_prose():
    """模型评价闸门：条数不足与非编号大段分别拦截。"""
    from tools.docx.core.structure_validation import _model_eval_bullet_format_issues

    doc = Document()
    for tx in ['七、模型评价与改进', '7.1 优点', '（1）参数较少',
               '7.2 局限与改进方向', '（1）忽略动态波动',
               '综合来看，本模型整体表现良好，具备较强实用性与推广价值，可满足工程应用要求，应进一步推广应用。']:
        doc.add_paragraph(tx)
    issues = _model_eval_bullet_format_issues(doc)
    assert any('优点仅 1 条' in i for i in issues)
    assert any('局限与改进仅 1 条' in i for i in issues)
    assert any('非编号正文段' in i for i in issues)



