"""命名契约与 SPSS 强制逻辑回归测试。

覆盖 2026-08 修复：
- append_code_files 默认 pattern 为 code/Q*.py（Q1.py 规范命名直接进附录；viz.py/build_paper/common 不进）
- load_spss_outputs 的 required=true 强制（漏填抛 ValueError）
- project_cleanup 的 build_paper 纯净性检查（不得 import solve_common / Q<序号>.py）
- _appendix_size_issues 对整题纯理论赛题豁免附录代码要求
"""
import json
import sys
import pathlib

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from tools.docx.core.paper_format import append_code_files, HEADING3_STYLE
from tools.docx.core.result_contract import load_spss_outputs
from tools.docx.core.structure_validation import _appendix_size_issues, _project_has_code
from tools.project_ops.project_cleanup import _check_build_paper_purity


def _mk_doc():
    doc = Document()
    try:
        st = doc.styles.add_style(HEADING3_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "黑体"
    except (KeyError, ValueError):
        pass
    return doc


def _spss_json(tmp_path, stats, tool="SPSS 27 手动"):
    datadir = tmp_path / "results" / "数据"
    datadir.mkdir(parents=True, exist_ok=True)
    (datadir / "spss_outputs.json").write_text(
        json.dumps({"tool": tool, "stats": stats}), encoding="utf-8"
    )


class TestAppendCodeFilesPattern:
    def test_only_solution_entry_rendered(self, tmp_path):
        code = tmp_path / "code"
        code.mkdir()
        (code / "Q1.py").write_text("# c\ndef solve():\n    return 1\n", encoding="utf-8")
        (code / "Q2.py").write_text("def f():\n    return 2\n", encoding="utf-8")
        (code / "Q1_求解.py").write_text("def g():\n    return 3\n", encoding="utf-8")  # 合法子模块
        (code / "viz.py").write_text("def plot():\n    pass\n", encoding="utf-8")  # 生图配置，不进附录
        (code / "build_paper.py").write_text("print('x')\n", encoding="utf-8")
        (code / "common.py").write_text("def u():\n    return 3\n", encoding="utf-8")

        doc = _mk_doc()
        append_code_files(doc, str(tmp_path))
        headings = [p.text for p in doc.paragraphs if p.text]
        assert 'Q1.py' in headings
        assert 'Q2.py' in headings
        assert any('Q1_求解.py' in h for h in headings)
        assert '附录A 支撑材料' in headings
        assert any('附录B' in h for h in headings)
        assert any('表' in h and '核心代码' in h for h in headings)
        assert not any('viz.py' in h for h in headings)
        assert not any('common.py' in h for h in headings)
        assert not any('build_paper.py' in h for h in headings)

    def test_manifest_source_scripts_preferred(self, tmp_path):
        code = tmp_path / "code"
        results = tmp_path / "results"
        code.mkdir()
        results.mkdir()
        (code / "Q1.py").write_text("# c\ndef solve():\n    return 1\n", encoding="utf-8")
        (code / "solve_common.py").write_text("def util():\n    return 42\n", encoding="utf-8")
        (results / "run_manifest.json").write_text(
            json.dumps({
                "schema_version": 1,
                "source_scripts": [
                    {"path": "code/Q1.py", "sha256": "x"},
                    {"path": "code/solve_common.py", "sha256": "x"},
                    {"path": "code/build_paper.py", "sha256": "x"},
                ],
            }),
            encoding="utf-8",
        )
        doc = _mk_doc()
        append_code_files(doc, str(tmp_path))
        headings = [p.text for p in doc.paragraphs if p.text]
        assert 'Q1.py' in headings
        assert 'solve_common.py' in headings
        assert '附录A 支撑材料' in headings
        assert any('附录B' in h for h in headings)


class TestSpssRequired:
    def test_required_missing_raises(self, tmp_path):
        _spss_json(tmp_path, [
            {"name": "配对t", "unit": "无量纲", "required": True},
            {"name": "d", "unit": "无量纲", "value": 0.42},
        ])
        with pytest.raises(ValueError, match="required"):
            load_spss_outputs(str(tmp_path))

    def test_optional_missing_skipped(self, tmp_path):
        _spss_json(tmp_path, [
            {"name": "可选项", "unit": "无量纲"},
            {"name": "F值", "unit": "无量纲", "value": 12.34},
        ])
        out = load_spss_outputs(str(tmp_path))
        assert [(e["name"], e["value"]) for e in out] == [("F值", 12.34)]
        assert out[0]["source"] == "results/数据/spss_outputs.json"

    def test_no_file_returns_empty(self, tmp_path):
        assert load_spss_outputs(str(tmp_path)) == []

    def test_zero_value_kept(self, tmp_path):
        _spss_json(tmp_path, [{"name": "p值", "unit": "无量纲", "value": 0.0}])
        out = load_spss_outputs(str(tmp_path))
        assert out[0]["value"] == 0.0


class TestBuildPaperPurity:
    def test_dependency_detected(self, tmp_path):
        code = tmp_path / "code"
        code.mkdir()
        (code / "build_paper.py").write_text(
            "from solve_common import load_csv\nimport Q1\nfrom tools.docx.core import paper_format\n",
            encoding="utf-8",
        )
        warnings = _check_build_paper_purity(tmp_path)
        assert len(warnings) == 1
        assert "solve_common" in warnings[0]

    def test_clean_build_paper(self, tmp_path):
        code = tmp_path / "code"
        code.mkdir()
        (code / "build_paper.py").write_text(
            "from tools.docx.core import paper_format\nfrom tools.common.io_utils import sha256_file\n",
            encoding="utf-8",
        )
        assert _check_build_paper_purity(tmp_path) == []

    def test_missing_build_paper(self, tmp_path):
        assert _check_build_paper_purity(tmp_path) == []


class TestAppendixPureTheoryExemption:
    def _mkdoc_with_appendix(self, paras=(), cells=()):
        doc = _mk_doc()
        try:
            doc.add_paragraph("附录", style="一级标题")
        except (KeyError, ValueError):
            doc.add_paragraph("附录")
        for t in paras:
            doc.add_paragraph(t)
        if cells:
            tb = doc.add_table(rows=1, cols=1)
            tb.rows[0].cells[0].text = cells[0]
        return doc

    def test_no_code_project_empty_appendix_ok(self, tmp_path):
        doc = self._mkdoc_with_appendix()
        assert _appendix_size_issues(doc, str(tmp_path)) == []

    def test_with_code_empty_appendix_fatal(self, tmp_path):
        (tmp_path / "code").mkdir()
        (tmp_path / "code" / "Q1.py").write_text("def solve():\n    return 1\n", encoding="utf-8")
        doc = self._mkdoc_with_appendix()
        assert _appendix_size_issues(doc, str(tmp_path)) != []

    def test_with_code_text_only_appendix_fatal(self, tmp_path):
        (tmp_path / "code").mkdir()
        (tmp_path / "code" / "Q1.py").write_text("def solve():\n    return 1\n", encoding="utf-8")
        doc = self._mkdoc_with_appendix(paras=["这里是附录文字说明"])
        assert _appendix_size_issues(doc, str(tmp_path)) != []

    def test_build_paper_not_counted_as_solution(self, tmp_path):
        code = tmp_path / "code"
        code.mkdir()
        (code / "build_paper.py").write_text("print('x')\n", encoding="utf-8")
        assert _project_has_code(str(tmp_path)) is False
