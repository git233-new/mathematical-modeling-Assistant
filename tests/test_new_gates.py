# -*- coding: utf-8 -*-
"""写作闸门的单元测试。

覆盖：模型假设复述拦截、摘要分问分段、一级标题顺序、题目一行、
表后段前一行、三线表自适应列宽、有效字号闸门、文本遮挡检测。
"""
import sys
import pathlib

import pytest
from docx.shared import Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from tools.docx.core import paper_format as pf
from tools.docx.core.paper_format import _validate_assumption_format
from tools.docx.core.structure_validation import (
    _abstract_paragraph_issues,
    _model_assumption_issues,
    _section_order_issues,
)


def _mk_doc():
    return pf.new_document()


def _add_assumptions(doc, texts):
    # 直接加段落：写入侧拦截由 test_assumption_write_time_guard 单独覆盖
    pf.heading1(doc, "三、模型假设")
    for t in texts:
        doc.add_paragraph(t, style=pf.BODY_STYLE)


# ── 1. 模型假设：复述题目拦截 ──

def test_assumption_restatement_rejected():
    doc = _mk_doc()
    texts = [f"假设 {i}（假设{i}）：所选模型采用的简化条件{i}。" for i in range(1, 6)]
    texts.append("假设 6（数据条件）：题目所给数据的测量误差服从正态分布。")
    _add_assumptions(doc, texts)
    issues = _model_assumption_issues(doc)
    assert any("复述题目规定条件" in i for i in issues)
    assert not any("依据：" in i for i in issues)  # 假设不含三链字样


def test_assumption_clean_passes_without_sanchain():
    doc = _mk_doc()
    texts = [f"假设 {i}（假设{i}）：所选算法采用的第{i}类简化，参数在合理区间内取值。" for i in range(1, 7)]
    _add_assumptions(doc, texts)
    assert _model_assumption_issues(doc) == []


def test_assumption_write_time_guard():
    with pytest.raises(ValueError, match="复述题目规定条件"):
        _validate_assumption_format("假设 1（数据）：题目给出的数据完整无缺失。")
    with pytest.raises(ValueError, match="过长"):
        _validate_assumption_format("假设 1（过长）：" + "细节" * 80)
    # 假设不写 依据/检验 字样
    assert _validate_assumption_format("假设 1（正态）：扰动项服从独立同分布正态噪声。") is None


# ── 2. 摘要分问分段 ──

_FILLER = "该问的求解流程、参数设置与数据来源均在正文对应章节交代，结果经稳健性检验支撑。"


def _abstract_doc(paras):
    # 摘要正文总量须过 400 字下限，按需补充中性说明句
    total = sum(len(t) for t in paras)
    while total < 420:
        paras = [t + _FILLER for t in paras]
        total = sum(len(t) for t in paras)
    doc = _mk_doc()
    pf.title(doc, "测试论文题目")
    pf.abstract_title(doc)
    for t in paras:
        pf.body(doc, t)
    pf.keywords(doc, "优化")
    return doc


QUANT = "求得最优目标值为 1234 元，相对误差为 3.2%。"


def test_abstract_two_questions_one_paragraph_rejected():
    doc = _abstract_doc([
        "本文针对某问题建立模型求解。" + QUANT,
        "针对问题一，建立模型A，求得结果为 100 元。" + QUANT + "针对问题二，建立模型B，求得结果为 200 元。",
    ])
    issues = _abstract_paragraph_issues(doc)
    assert any("按问题分段" in i for i in issues)


def test_abstract_per_question_paragraphs_pass():
    doc = _abstract_doc([
        "本文针对某问题建立模型求解。" + QUANT,
        "针对问题一，建立模型A，求得结果为 100 元，误差 3%。" + QUANT,
        "针对问题二，建立模型B，求得结果为 200 元，误差 4%。" + QUANT,
        "针对问题三，建立模型C，求得结果为 300 元，误差 5%。" + QUANT,
    ])
    assert _abstract_paragraph_issues(doc) == []


def test_abstract_three_questions_one_paragraph_rejected():
    doc = _abstract_doc([
        "本文针对某问题建立模型求解。" + QUANT,
        "针对问题一，建立模型A，求得结果为 100 元，误差 3%。" + QUANT
        + "针对问题二，建立模型B，求得结果为 200 元，误差 4%。针对问题三，建立模型C，求得结果为 300 元。",
    ])
    issues = _abstract_paragraph_issues(doc)
    assert any("按问题分段" in i for i in issues)


def test_abstract_without_quantified_result_rejected():
    doc = _abstract_doc([
        "本文针对某问题建立了较好的模型，效果良好。",
        "针对问题一，模型A表现优秀。",
    ])
    assert any("量化结果" in i for i in _abstract_paragraph_issues(doc))


# ── 3. 一级标题顺序：AI声明 → 参考文献 ──

def test_section_order_correct_passes():
    doc = _mk_doc()
    pf.heading1(doc, "AI工具使用声明")
    pf.heading1(doc, "参考文献")
    assert _section_order_issues(doc) == []


def test_section_order_ai_after_references_rejected():
    doc = _mk_doc()
    pf.heading1(doc, "参考文献")
    pf.heading1(doc, "AI工具使用声明")
    issues = _section_order_issues(doc)
    assert any("AI工具使用声明必须在参考文献之前" in i for i in issues)


def test_section_order_alias_and_missing_sections_pass():
    doc = _mk_doc()
    pf.heading1(doc, "AI工具使用详情")
    pf.heading1(doc, "参考文献")
    assert _section_order_issues(doc) == []
    doc2 = _mk_doc()
    doc2.add_paragraph("一、问题重述", style=pf.HEADING1_STYLE)
    assert _section_order_issues(doc2) == []


# ── 4. 论文题目一行（≤24 字） ──

def test_title_over_24_chars_rejected():
    doc = _mk_doc()
    with pytest.raises(ValueError, match="超过 24 字"):
        pf.title(doc, "基于超长题目的某某某模型构建与某地区某某问题的某某优化研究超出限制")


def test_title_24_chars_ok_and_formatted():
    doc = _mk_doc()
    text = "基于组合模型的某地区用水量预测与调度优化研究"
    assert len(text) == 22
    p = pf.title(doc, text)
    assert p.alignment is not None
    assert p.paragraph_format.first_line_indent == Pt(0)


# ── 5. 表后正文段前一行 ──

def test_body_after_table_gets_space_before():
    doc = _mk_doc()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "方法"
    p_after_table = pf.body(doc, "表格下方的解释文字。")
    assert p_after_table.paragraph_format.space_before == Pt(15)
    p_normal = pf.body(doc, "普通正文段落没有段前空行。")
    assert p_normal.paragraph_format.space_before is None or p_normal.paragraph_format.space_before == Pt(0)


# ── 6. 三线表自适应列宽 ──

def _twips(col):
    from docx.oxml.ns import qn
    return int(col.cells[0]._tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w')))


def test_three_line_widths_equal_split_fills_body_width():
    """三线表列宽为等分铺满版心。"""
    doc = _mk_doc()
    table = doc.add_table(rows=3, cols=4)
    for r in range(3):
        for c in range(4):
            table.cell(r, c).text = f"内容{r}{c}"
    pf._assign_three_line_widths(table, doc)
    widths = [_twips(c) for c in table.columns]
    assert len(set(widths)) == 1             # 各列等宽
    assert abs(sum(widths) - 8300) <= 2      # 总宽精确铺满版心
    from docx.oxml.ns import qn
    grid = table._tbl.find(qn('w:tblGrid'))
    assert len(grid.findall(qn('w:gridCol'))) == 4


# ── 7. 图件导出闸门：有效字号 + 文本遮挡 ──

def test_effective_font_size_gate():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    from tools.figure.templates.mm_style import configure_matplotlib, effective_font_pt, ensure_print_font_size

    configure_matplotlib()  # 基线字号 11pt（模板出图前的标准动作）

    style = plt.figure(figsize=(13.8, 4.6))
    with pytest.raises(ValueError, match="有效字号"):
        ensure_print_font_size(style, print_width_cm=12)
    plt.close(style)

    ok = plt.figure(figsize=(4.7, 3.0))
    eff = ensure_print_font_size(ok, print_width_cm=12)
    assert eff >= 9.0
    assert abs(effective_font_pt(ok, print_width_cm=12) - 11.0) < 0.3
    plt.close(ok)


def test_text_overlap_detection():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    from tools.figure.runtime.figure_safety import text_overlap_issues

    fig, ax = plt.subplots(figsize=(4.7, 3.0))
    ax.text(0.5, 0.5, "重叠文本甲", ha="center")
    ax.text(0.5, 0.5, "重叠文本乙", ha="center")
    assert text_overlap_issues(fig)
    plt.close(fig)

    fig2, ax2 = plt.subplots(figsize=(4.7, 3.0))
    ax2.text(0.2, 0.8, "分离文本甲")
    ax2.text(0.8, 0.2, "分离文本乙")
    assert not text_overlap_issues(fig2)
    plt.close(fig2)

def test_symbol_orphan_symbols_rejected(tmp_path):
    """符号表符号必须在正文数学公式中出现：公式未用的 β、全程未用的 γ 都拒存。"""
    from docx.oxml.ns import qn
    from tools.docx.core.structure_validation import _symbol_orphan_issues

    doc = _mk_doc()
    pf.heading1(doc, "三、模型假设")
    pf.body(doc, "本文采用模型求解，加速度 α 与阻尼系数 β 由实验数据拟合得到。" + "该问的求解流程与参数设置均在正文交代。" * 8)
    pf.heading1(doc, "四、符号说明")
    table = doc.add_table(rows=4, cols=3)
    for row, cells in zip(table.rows, (
        ("符号", "说明", "单位"),
        ("α", "加速度", "m/s2"),
        ("β", "阻尼系数", "-"),
        ("γ", "全程未用的摆设符号", "-"),
    )):
        for cell, t in zip(row.cells, cells):
            cell.text = t
    pf.heading1(doc, "五、模型建立与求解")
    p = doc.add_paragraph()
    omml = p._p.makeelement(qn('m:oMath'), {})
    mt = omml.makeelement(qn('m:t'), {})
    mt.text = "α"
    omml.append(mt)
    p._p.append(omml)
    pf.body(doc, "模型中摩擦系数 μ 与加速度共同决定响应特性，推导详见公式。" + "该章说明。" * 6)
    issues = _symbol_orphan_issues(doc)
    assert any("γ" in i and "公式" in i for i in issues)
    assert any("β" in i and "公式" in i for i in issues)
    assert not any("α" in i for i in issues)


def test_symbol_used_in_formula_passes(tmp_path):
    """符号仅在公式（OMML）中出现也算已使用 → 不预警。"""
    from docx.oxml.ns import qn
    from tools.docx.core.structure_validation import _symbol_orphan_issues

    doc = _mk_doc()
    pf.heading1(doc, "四、符号说明")
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "符号"
    table.rows[1].cells[0].text = "δ"
    pf.heading1(doc, "五、模型建立与求解")
    p = doc.add_paragraph()
    omml = p._p.get_or_add_pPr().makeelement(qn('m:oMath'), {})
    mt = omml.makeelement(qn('m:t'), {})
    mt.text = "δ"
    omml.append(mt)
    p._p.append(omml)
    pf.body(doc, "公式的推导说明文字，解释该符号的含义与取值来源。" * 3)
    assert _symbol_orphan_issues(doc) == []
